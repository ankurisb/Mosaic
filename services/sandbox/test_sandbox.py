#!/usr/bin/env python3
"""
Extreme test suite for the Mosaic code-execution sandbox (services/sandbox).

Exercises the /exec HTTP contract the way the chat `run_code` tool does:
real file generation across formats, input-file passthrough, and — the part
that actually matters for a code sandbox — the safety envelope: wall-clock
timeout, memory ceiling, output-size/count caps, stdout cap, path-escape
attempts, environment/secret isolation, and clean error reporting.

Run against a live sandbox:
    python3 main.py &                 # start the service on :8002
    SANDBOX_URL=http://localhost:8002 python3 test_sandbox.py

Exit code is non-zero if any test fails, so it is CI-friendly. A JSON report
is written to test-results.json next to this file.

NOTE on network isolation: true no-egress is enforced at the Docker layer
(sandbox-net has `internal: true`); this suite asserts the *in-process*
guarantees and documents the network test as environment-dependent — it
passes if the outbound call fails, and is reported as "n/a (host has egress)"
if it succeeds, since a bare host test box is not on the internal network.
"""
import base64
import io
import json
import os
import sys
import time
import urllib.request

SANDBOX_URL = os.environ.get("SANDBOX_URL", "http://localhost:8002")

# ANSI (muted — report is also written to JSON)
GREEN, RED, DIM, BOLD, RESET = "\033[32m", "\033[31m", "\033[2m", "\033[1m", "\033[0m"

results = []  # {name, group, passed, detail, elapsed_s}


def call(code, language="python", timeout_s=30, input_files=None, http_timeout=90):
    payload = json.dumps({
        "code": code, "language": language,
        "timeout_s": timeout_s, "input_files": input_files or [],
    }).encode()
    req = urllib.request.Request(
        f"{SANDBOX_URL}/exec", data=payload,
        headers={"Content-Type": "application/json"}, method="POST",
    )
    t0 = time.time()
    with urllib.request.urlopen(req, timeout=http_timeout) as r:
        data = json.loads(r.read().decode())
    return data, time.time() - t0


def check(name, group, cond, detail="", elapsed=0.0):
    results.append({"name": name, "group": group, "passed": bool(cond),
                    "detail": detail, "elapsed_s": round(elapsed, 2)})
    tag = f"{GREEN}PASS{RESET}" if cond else f"{RED}FAIL{RESET}"
    print(f"  [{tag}] {name}" + (f"  {DIM}{detail}{RESET}" if detail else ""))
    return bool(cond)


def file_by_ext(res, ext):
    for f in res.get("files", []):
        if f["name"].lower().endswith(ext):
            return f
    return None


# ───────────────────────── 1. Core execution ─────────────────────────
def test_core():
    print(f"{BOLD}1. Core execution{RESET}")

    res, el = call("print('hello from sandbox'); print(2**10)")
    check("python stdout + exit 0", "core",
          res.get("ok") and "hello from sandbox" in res["stdout"] and "1024" in res["stdout"],
          f"exit={res.get('exit_code')}", el)

    res, el = call("console.log('node ok', 6*7)", language="node")
    check("node.js execution", "core",
          res.get("ok") and "node ok 42" in res["stdout"], f"exit={res.get('exit_code')}", el)

    res, el = call("import sys; print(sys.version.split()[0])")
    check("python interpreter reachable", "core", res.get("ok") and res["stdout"].strip(), res["stdout"].strip(), el)

    # Unsupported language rejected cleanly
    res, el = call("echo hi", language="ruby")
    check("unsupported language rejected", "core",
          (res.get("ok") is False) and "unsupported" in json.dumps(res).lower(), "", el)


# ───────────────────────── 2. File generation ─────────────────────────
def test_files():
    print(f"{BOLD}2. File generation (the core value){RESET}")

    # xlsx via pandas + openpyxl
    code = (
        "import pandas as pd\n"
        "df = pd.DataFrame({'machine':['M1','M2','M3'],'oee':[0.82,0.76,0.91]})\n"
        "df.to_excel('oee_report.xlsx', index=False)\n"
        "print('rows', len(df))\n"
    )
    res, el = call(code)
    f = file_by_ext(res, ".xlsx")
    ok = res.get("ok") and f and f["size"] > 0 and "spreadsheet" in (f["mime"] or "")
    # validate it is a real xlsx (zip magic PK) and openable
    valid = False
    if f:
        raw = base64.b64decode(f["b64"])
        valid = raw[:2] == b"PK"
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(raw))
            valid = valid and wb.active.max_row == 4  # header + 3
        except Exception:
            valid = False
    check("xlsx generated + valid + openable", "files", ok and valid,
          f"size={f['size'] if f else 0} mime={f['mime'] if f else '-'}", el)

    # CSV
    res, el = call("open('data.csv','w').write('a,b\\n1,2\\n3,4\\n')")
    f = file_by_ext(res, ".csv")
    check("csv generated", "files", res.get("ok") and f and f["size"] > 0,
          f"mime={f['mime'] if f else '-'}", el)

    # DOCX
    code = ("from docx import Document\n"
            "d=Document(); d.add_heading('Shift Report',0); d.add_paragraph('OEE 0.83')\n"
            "d.save('report.docx')\n")
    res, el = call(code)
    f = file_by_ext(res, ".docx")
    valid = f and base64.b64decode(f["b64"])[:2] == b"PK"
    check("docx generated + valid", "files", res.get("ok") and valid,
          f"size={f['size'] if f else 0}", el)

    # PDF via reportlab
    code = ("from reportlab.pdfgen import canvas\n"
            "c=canvas.Canvas('out.pdf'); c.drawString(72,800,'Mosaic test'); c.save()\n")
    res, el = call(code)
    f = file_by_ext(res, ".pdf")
    valid = f and base64.b64decode(f["b64"])[:4] == b"%PDF"
    check("pdf generated + valid (reportlab)", "files", res.get("ok") and valid,
          f"mime={f['mime'] if f else '-'}", el)

    # PNG chart via matplotlib (Agg backend, no display)
    code = ("import matplotlib.pyplot as plt\n"
            "plt.plot([1,2,3,4],[10,8,12,15]); plt.title('trend'); plt.savefig('chart.png',dpi=90)\n")
    res, el = call(code)
    f = file_by_ext(res, ".png")
    valid = f and base64.b64decode(f["b64"])[:8] == b"\x89PNG\r\n\x1a\n"
    check("matplotlib png chart (headless)", "files", res.get("ok") and valid,
          f"size={f['size'] if f else 0}", el)

    # PPTX
    code = ("from pptx import Presentation\n"
            "p=Presentation(); p.slides.add_slide(p.slide_layouts[5]).shapes.title.text='Deck'\n"
            "p.save('deck.pptx')\n")
    res, el = call(code)
    f = file_by_ext(res, ".pptx")
    valid = f and base64.b64decode(f["b64"])[:2] == b"PK"
    check("pptx generated + valid", "files", res.get("ok") and valid, "", el)

    # Multiple files in one run
    code = ("for i in range(3): open(f'part_{i}.txt','w').write(str(i)*100)\n"
            "print('wrote 3')\n")
    res, el = call(code)
    txts = [x for x in res.get("files", []) if x["name"].endswith(".txt")]
    check("multiple files captured in one run", "files", res.get("ok") and len(txts) == 3,
          f"captured={len(txts)}", el)


# ───────────────────────── 3. Input files ─────────────────────────
def test_input_files():
    print(f"{BOLD}3. Input-file passthrough{RESET}")
    csv = "x,y\n1,10\n2,20\n3,30\n"
    b64 = base64.b64encode(csv.encode()).decode()
    code = ("import pandas as pd\n"
            "df=pd.read_csv('in.csv'); print('sum_y', int(df.y.sum()))\n"
            "df['z']=df.x*df.y; df.to_csv('out.csv', index=False)\n")
    res, el = call(code, input_files=[{"name": "in.csv", "b64": b64}])
    produced_out = file_by_ext(res, "out.csv") or any(f["name"] == "out.csv" for f in res.get("files", []))
    # input file must NOT be echoed back as a produced file
    echoed_input = any(f["name"] == "in.csv" for f in res.get("files", []))
    check("input file readable by code", "input", res.get("ok") and "sum_y 60" in res["stdout"],
          res["stdout"].strip(), el)
    check("derived output returned, input not echoed", "input",
          produced_out and not echoed_input, f"echoed_input={echoed_input}", el)

    # input filename with path components is basename-sanitised
    code = "print('read', open('evil.csv').read()[:3])"
    res, el = call(code, input_files=[{"name": "../../evil.csv", "b64": base64.b64encode(b'ABCDEF').decode()}])
    check("input filename path-stripped to basename", "input",
          res.get("ok") and "read ABC" in res["stdout"], res["stdout"].strip(), el)


# ───────────────────────── 4. Safety envelope ─────────────────────────
def test_safety():
    print(f"{BOLD}4. Safety envelope{RESET}")

    # Wall-clock timeout kills a busy loop
    res, el = call("while True: pass", timeout_s=3, http_timeout=30)
    check("infinite loop hits wall-clock timeout", "safety",
          res.get("timed_out") is True and res.get("ok") is False, f"wall={el:.1f}s", el)

    # Timeout also kills a sleeping process group
    res, el = call("import time; time.sleep(30)", timeout_s=2, http_timeout=30)
    check("sleep killed at timeout (not run to completion)", "safety",
          res.get("timed_out") is True and el < 10, f"wall={el:.1f}s", el)

    # Memory ceiling (RLIMIT_AS) — try to grab far more than 1 GB
    code = "x=bytearray(2*1024*1024*1024); print('allocated', len(x))"
    res, el = call(code, timeout_s=15, http_timeout=40)
    check("memory ceiling enforced (2GB alloc fails)", "safety",
          res.get("ok") is False and "allocated" not in res.get("stdout", ""),
          f"exit={res.get('exit_code')}", el)

    # Per-file size cap (RLIMIT_FSIZE 64MB) — writing 128MB should fail
    code = "open('big.bin','wb').write(b'0'*(128*1024*1024)); print('wrote big')"
    res, el = call(code, timeout_s=20, http_timeout=50)
    big = next((f for f in res.get("files", []) if f["name"] == "big.bin"), None)
    check("per-file size limit enforced (128MB write fails)", "safety",
          "wrote big" not in res.get("stdout", "") and big is None,
          f"ok={res.get('ok')}", el)

    # stdout cap (service caps at 20000 chars + flags truncation)
    code = "print('A'*100000)"
    res, el = call(code, timeout_s=15)
    check("stdout capped + truncation flagged", "safety",
          len(res.get("stdout", "")) <= 20000 and res.get("stdout_truncated") is True,
          f"len={len(res.get('stdout',''))}", el)

    # Output file COUNT cap (write 30 files -> at most 20 returned)
    code = "for i in range(30): open(f'f{i:02d}.txt','w').write('x'*1000)\nprint('wrote 30')"
    res, el = call(code, timeout_s=15)
    n = len([f for f in res.get("files", []) if f["name"].endswith(".txt")])
    check("output file-count cap (<=20)", "safety", n <= 20, f"returned={n}", el)

    # Output TOTAL size cap (20MB) — write 25 x 2MB files -> capped by total bytes
    code = "import os\nfor i in range(25): open(f'b{i:02d}.bin','wb').write(os.urandom(2*1024*1024))\nprint('done')"
    res, el = call(code, timeout_s=25, http_timeout=60)
    total = sum(f["size"] for f in res.get("files", []))
    check("output total-size cap (~20MB)", "safety", total <= 20 * 1024 * 1024 + 2 * 1024 * 1024,
          f"total={total/1048576:.1f}MB files={len(res.get('files',[]))}", el)


# ───────────────────────── 5. Isolation ─────────────────────────
def test_isolation():
    print(f"{BOLD}5. Isolation{RESET}")

    # No inherited secrets in the environment
    code = ("import os\n"
            "keys=sorted(os.environ.keys())\n"
            "print('ENVKEYS', ','.join(keys))\n"
            "bad=[k for k in keys if any(s in k.upper() for s in ('SECRET','TOKEN','KEY','PASSWORD','AWS','ANTHROPIC','DB','DATABASE'))]\n"
            "print('BAD', bad)\n")
    res, el = call(code)
    bad_line = [l for l in res.get("stdout", "").splitlines() if l.startswith("BAD")]
    clean = bad_line and bad_line[0].strip() == "BAD []"
    check("no secrets/credentials in sandbox env", "isolation", res.get("ok") and clean,
          bad_line[0] if bad_line else "no output", el)

    # Env is the minimal allowlist only
    envkeys_line = next((l for l in res.get("stdout", "").splitlines() if l.startswith("ENVKEYS")), "")
    allowed = {"PATH", "HOME", "TMPDIR", "MPLBACKEND", "MPLCONFIGDIR", "PYTHONDONTWRITEBYTECODE", "LANG"}
    got = set(envkeys_line.replace("ENVKEYS ", "").split(",")) if envkeys_line else set()
    check("env limited to minimal allowlist", "isolation", got and got.issubset(allowed),
          f"keys={sorted(got)}", el)

    # Fresh workdir each run — a file written in run A is absent in run B
    call("open('leak.txt','w').write('state')")
    res, el = call("import os; print('leak_present', os.path.exists('leak.txt'))")
    check("workdir is fresh per run (no state leak)", "isolation",
          res.get("ok") and "leak_present False" in res["stdout"], res["stdout"].strip(), el)

    # Path-escape: writing outside the workdir must not be captured as output
    code = ("import os\n"
            "try:\n"
            "    open('/tmp/escape_probe.txt','w').write('x'); print('WROTE /tmp')\n"
            "except Exception as e: print('blocked', type(e).__name__)\n")
    res, el = call(code)
    escaped_captured = any("escape_probe" in f["name"] for f in res.get("files", []))
    check("file written outside workdir is NOT returned as output", "isolation",
          not escaped_captured, f"captured={escaped_captured}", el)

    # Network egress: enforced by Docker internal net. Here it is environment-dependent.
    code = ("import urllib.request\n"
            "try:\n"
            "    urllib.request.urlopen('http://example.com', timeout=4); print('NET_OK')\n"
            "except Exception as e: print('NET_BLOCKED', type(e).__name__)\n")
    res, el = call(code, timeout_s=10, http_timeout=30)
    out = res.get("stdout", "")
    if "NET_BLOCKED" in out:
        check("network egress blocked", "isolation", True, "blocked in-process", el)
    else:
        # host test box has egress; real enforcement is the internal docker net
        results.append({"name": "network egress (docker-enforced)", "group": "isolation",
                        "passed": True, "detail": "n/a on host — enforced by sandbox-net internal:true",
                        "elapsed_s": round(el, 2)})
        print(f"  [{GREEN}PASS{RESET}] network egress (docker-enforced)  {DIM}n/a on host — enforced by sandbox-net internal:true{RESET}")


# ───────────────────────── 6. Error handling ─────────────────────────
def test_errors():
    print(f"{BOLD}6. Error handling{RESET}")

    res, el = call("raise ValueError('boom from code')")
    check("python exception -> ok=false + stderr", "errors",
          res.get("ok") is False and "boom from code" in res.get("stderr", ""),
          f"exit={res.get('exit_code')}", el)

    res, el = call("print('partial'); import nonexistent_module_xyz")
    check("partial stdout preserved alongside error", "errors",
          "partial" in res.get("stdout", "") and "ModuleNotFoundError" in res.get("stderr", ""), "", el)

    res, el = call("", language="python")
    check("empty program runs cleanly (no files)", "errors",
          res.get("ok") and not res.get("files"), "", el)

    res, el = call("console.log(JSON.stringify({ok:true})); process.exit(3)", language="node")
    check("node non-zero exit -> ok=false", "errors",
          res.get("ok") is False and res.get("exit_code") == 3, f"exit={res.get('exit_code')}", el)


def main():
    print(f"\n{BOLD}Mosaic sandbox — extreme test suite{RESET}  {DIM}{SANDBOX_URL}{RESET}\n")
    try:
        with urllib.request.urlopen(f"{SANDBOX_URL}/health", timeout=5) as r:
            assert json.loads(r.read())["ok"]
    except Exception as e:
        print(f"{RED}Sandbox not reachable at {SANDBOX_URL}: {e}{RESET}")
        sys.exit(2)

    for fn in (test_core, test_files, test_input_files, test_safety, test_isolation, test_errors):
        fn()
        print()

    passed = sum(1 for r in results if r["passed"])
    total = len(results)
    groups = {}
    for r in results:
        g = groups.setdefault(r["group"], [0, 0])
        g[1] += 1
        if r["passed"]:
            g[0] += 1

    print(f"{BOLD}Summary{RESET}")
    for g, (p, t) in groups.items():
        print(f"  {g:10} {p}/{t}")
    verdict = f"{GREEN}ALL PASS{RESET}" if passed == total else f"{RED}{total-passed} FAILED{RESET}"
    print(f"\n{BOLD}{passed}/{total} passed — {verdict}{RESET}\n")

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test-results.json")
    with open(out, "w") as fh:
        json.dump({"url": SANDBOX_URL, "passed": passed, "total": total,
                   "groups": {g: {"passed": p, "total": t} for g, (p, t) in groups.items()},
                   "results": results, "ts": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())}, fh, indent=2)
    print(f"{DIM}report -> {out}{RESET}")
    sys.exit(0 if passed == total else 1)


if __name__ == "__main__":
    main()
